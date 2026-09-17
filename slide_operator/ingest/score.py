"""Read the printed music on melody slides (optical music recognition).

Notes and lyrics answer different questions, and each is blind where the other
sees:

  * **Notes say where you are within the tune.** Matched against the audio they
    also tell right music from wrong — measured DTW cost 0.389 against the
    canticle's own audio, 0.562 against a different hymn, 0.701 against speech.
  * **Notes cannot say which stanza.** A strophic hymn prints the same tune on
    every slide: all four slides of hymn 733 produced identical note sequences.
  * **Lyrics say which stanza**, but on sung audio their alignment scores barely
    separate right text from wrong (0.031 against 0.023), so they cannot be
    trusted to place a boundary on their own.

So the two are used together, never alone.
"""
from __future__ import annotations

import io
import json
import subprocess
import tempfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .deck import MELODY_MIN_ASPECT, Slide


def _strips(pptx: Path, index: int) -> list[Image.Image]:
    slide = Presentation(str(pptx)).slides[index - 1]
    out = []
    for sh in sorted((s for s in slide.shapes
                      if s.shape_type == MSO_SHAPE_TYPE.PICTURE and s.height
                      and s.width / s.height >= MELODY_MIN_ASPECT),
                     key=lambda s: s.top):
        im = Image.open(io.BytesIO(sh.image.blob)).convert("RGBA")
        page = Image.new("RGB", im.size, "white")       # OMR wants ink on white
        page.paste(im, mask=im.getchannel("A"))
        out.append(page)
    return out


def _read_strip(img: Image.Image) -> list[int]:
    """Pitch classes of one staff, via homr. Empty if it cannot be read."""
    from music21 import converter

    with tempfile.TemporaryDirectory() as tmp:
        p = Path(tmp) / "staff.png"
        img.save(p)
        try:
            subprocess.run(["homr", str(p)], capture_output=True, timeout=180)
        except Exception:
            return []
        xml = list(Path(tmp).glob("*.musicxml"))
        if not xml:
            return []
        try:
            return [n.pitch.pitchClass for n in converter.parse(str(xml[0])).recurse().notes]
        except Exception:
            return []


def read(run: Path, slides: list[Slide], only: set[int] | None = None) -> dict[int, list[int]]:
    """Notated pitch classes per slide. Cached — OMR is slow and a deck is fixed."""
    cache = run / "notes.json"
    have = json.loads(cache.read_text()) if cache.exists() else {}
    wanted = [s for s in slides
              if s.melody_lines and (only is None or s.index in only)
              and str(s.index) not in have]
    for s in wanted:
        pitches: list[int] = []
        for img in _strips(run / "deck.pptx", s.index):
            pitches += _read_strip(img)
        have[str(s.index)] = pitches
    if wanted:
        cache.write_text(json.dumps(have, indent=2))
    return {int(k): v for k, v in have.items()}
