"""Read lyrics out of sheet-music images on melody slides (macOS Vision OCR).

Some decks set hymns as notation, so a slide's sung words exist only as pictures.
Without them the slide has nothing to listen for. OCR runs before the service and
is cached per run, since the deck cannot change once the service starts.
"""
from __future__ import annotations

import io
import json
import re
from pathlib import Path

from PIL import Image, ImageOps
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .deck import MELODY_MIN_ASPECT, Slide

MIN_CONFIDENCE = 0.5
# Lyrics sit in a band under the staff. Handing the recognizer the staff lines and
# slurs as well makes it drop or corrupt words — measured on one canticle slide as
# a lost "who take" and a "faith-ful-ness" read as "faith-fu-ness". Cropping to the
# band fixes both. Upscaling does not help and introduces its own errors.
LYRIC_BAND = 0.42
_HYPHEN = re.compile(r"(\w)\s*[-–]\s*(\w)")     # "Sav - ior" -> "Savior"


def _ocr(blob: bytes) -> str:
    """Read a strip twice and keep the better reading.

    The staff lines and slurs make the recognizer drop or corrupt lyrics, so a
    crop to the lyric band often reads better — but on strips where the band clips
    the text it reads worse. Neither wins everywhere, so run both and keep the one
    that recovers more words.
    """
    full = _ocr_image(blob, 1.0)
    band = _ocr_image(blob, LYRIC_BAND)
    count = lambda t: sum(1 for x in t.split() if len(x) > 1 and any(c.isalpha() for c in x))
    return band if count(band) > count(full) else full


def _ocr_image(blob: bytes, band: float) -> str:
    import Vision
    from Foundation import NSData

    im = Image.open(io.BytesIO(blob)).convert("RGBA")
    alpha = im.getchannel("A")
    if alpha.getextrema()[0] < 255:
        ink = Image.eval(alpha, lambda a: 255 - a)   # faint glyphs on transparency
    else:
        ink = ImageOps.autocontrast(im.convert("L"))
    buf = io.BytesIO()
    if band < 1.0:
        w, h = ink.size
        ink = ink.crop((0, int(h * (1 - band)), w, h))
    ink.save(buf, "PNG")
    data = NSData.dataWithBytes_length_(buf.getvalue(), len(buf.getvalue()))
    handler = Vision.VNImageRequestHandler.alloc().initWithData_options_(data, None)
    req = Vision.VNRecognizeTextRequest.alloc().init()
    req.setRecognitionLevel_(Vision.VNRequestTextRecognitionLevelAccurate)
    req.setUsesLanguageCorrection_(False)
    handler.performRequests_error_([req], None)

    # Vision returns one box per word group; rebuild lines top to bottom, left to right.
    lines: list[dict] = []
    for o in req.results():
        c = o.topCandidates_(1)[0]
        if c.confidence() < MIN_CONFIDENCE or not re.search("[A-Za-z]", c.string()):
            continue
        bb = o.boundingBox()
        cy = bb.origin.y + bb.size.height / 2
        for line in lines:
            if abs(line["cy"] - cy) < bb.size.height * 0.6:
                line["items"].append((bb.origin.x, c.string()))
                break
        else:
            lines.append({"cy": cy, "items": [(bb.origin.x, c.string())]})
    lines.sort(key=lambda l: -l["cy"])
    text = " ".join(" ".join(s for _, s in sorted(l["items"])) for l in lines)
    return _HYPHEN.sub(r"\1\2", text)


def attach(slides: list[Slide], pptx_path: Path, cache: Path) -> None:
    if cache.exists():
        found = json.loads(cache.read_text())
    else:
        found = {}
        for s, raw in zip(slides, Presentation(str(pptx_path)).slides):
            if not s.melody_lines:
                continue
            pics = sorted(
                (sh for sh in raw.shapes
                 if sh.shape_type == MSO_SHAPE_TYPE.PICTURE and sh.height
                 and sh.width / sh.height >= MELODY_MIN_ASPECT),
                key=lambda sh: sh.top)
            found[str(s.index)] = " ".join(_ocr(sh.image.blob) for sh in pics)
        cache.write_text(json.dumps(found, indent=2))
    for s in slides:
        s.lyrics = found.get(str(s.index), "")
