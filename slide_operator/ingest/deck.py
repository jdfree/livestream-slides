"""Slide deck ingestion."""
from __future__ import annotations

from dataclasses import dataclass, asdict, field
from pathlib import Path

from .text import normalize, content_words, is_credit, names_a_tune


@dataclass
class Slide:
    index: int              # 1-based, as the operating doc numbers slides
    title: str              # title placeholder, if any
    body: str               # everything else, in shape order
    norm: str               # normalized body
    norm_title: str
    n_words: int            # content words in the body
    melody_lines: int = 0   # wide picture strips: sheet music with lyrics as images
    is_cover: bool = False  # displays nothing the congregation reads (3)
    lyrics: str = ""         # OCR of melody images, attached by ingest.melody
    run_id: int = 0          # consecutive slides sharing a title: one hymn, one reading
    musical: bool = False    # this element is sung (5.1, stanza clock)
    notes: list[str] = field(default_factory=list)


# A slide with almost no body text displays nothing anyone reads from. The
# threshold is deliberately generous: a title card like "Sermon" or
# "Distribution" has a handful of words at most, while the shortest real
# content slide in a liturgy runs to a dozen or more.
COVER_MAX_BODY_WORDS = 6
# Some decks set hymns as sheet music, so the lyrics are pictures and the slide
# has no body text at all. A melody line is a wide, short strip; a logo on a
# cover slide is not. Aspect is width/height in real proportions.
MELODY_MIN_ASPECT = 3.0


def load(path: str | Path) -> list[Slide]:
    from pptx import Presentation

    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(str(path))
    slides: list[Slide] = []
    for i, s in enumerate(prs.slides, start=1):
        melody = sum(
            1 for sh in s.shapes
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE and sh.height
            and sh.width / sh.height >= MELODY_MIN_ASPECT
        )
        title = ""
        body_parts: list[str] = []
        for shape in s.shapes:
            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text.strip()
            if not text:
                continue
            title_shape = s.shapes.title
            is_title = title_shape is not None and shape.shape_id == title_shape.shape_id
            if is_title and not title:
                title = text
            else:
                body_parts.append(text)

        body = "\n".join(body_parts)
        # Some decks put the whole slide in one box with no title placeholder.
        if not title and body_parts:
            first = body_parts[0].split("\n")[0]
            if len(first) < 60:
                title = first
        # The title is displayed, not spoken. Some decks repeat it at the top of
        # the body box; drop those lines so they never become listening anchors.
        if title:
            lines = body.split("\n")
            while lines and normalize(lines[0]) in ("", normalize(title)):
                lines.pop(0)
            body = "\n".join(lines).strip()

        words = content_words(body)
        slides.append(Slide(
            index=i,
            title=title,
            body=body,
            norm=normalize(body),
            norm_title=normalize(title),
            n_words=len(words),
            melody_lines=melody,
            is_cover=len(words) <= COVER_MAX_BODY_WORDS and melody == 0,
        ))
    mark_runs(slides)
    return slides


def to_json(slides: list[Slide]) -> list[dict]:
    return [asdict(s) for s in slides]


def mark_runs(slides: list[Slide]) -> None:
    """Group consecutive slides sharing a title, and mark sung runs.

    Sung is inferred from sheet music or from a "Tune:" credit — both typographic,
    neither a liturgical word list (2.2).
    """
    run = 0
    for i, s in enumerate(slides):
        if i and normalize(s.title) != normalize(slides[i - 1].title):
            run += 1
        s.run_id = run
    for r in {s.run_id for s in slides}:
        members = [s for s in slides if s.run_id == r]
        if any(s.melody_lines or names_a_tune(s.body) for s in members):
            for s in members:
                s.musical = True
